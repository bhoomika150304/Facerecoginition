import os, sys
from chardet import detect
import cv2
import numpy as np
import pickle
import random
import time
from datetime import datetime
import mail
import glob
from mail import*
from PIL import Image
import shutil
from openpyxl import load_workbook
import pandas as pd
import serial 


capture_duration = 30


faceDetect=cv2.CascadeClassifier('haarcascade_frontalface_default.xml');
cam=cv2.VideoCapture(0);
rec=cv2.face.LBPHFaceRecognizer_create();
rec.read("recognizer//model.yml")

net = cv2.dnn.readNetFromDarknet("yolo_custom.cfg","yolo_custom_last.weights");

all_attendance = []

def sent_mail(name_list, present_list):
    for i in range(len(name_list)):
        if name_list[i] in present_list:
            all_attendance.append('Present')
        else:
            all_attendance.append('Absent')
            
    return all_attendance

list_name = []
list_num = []
list_time = []

def writter(username, roll_number):
    if username not in list_name:
        print(username)
        list_name.append(username)
        list_num.append(roll_number)
        now = datetime.now()
        datestring = now.strftime("%H:%M:%S")
        list_time.append(datestring)
    

id=0
count = 0
count_ = 0
ret,img=cam.read();
start_time = time.time()
ids = '$,'


name_list_ = []
roll_num = []

with open('names.txt') as f:
    for line in f:
        name_list_.append(line)
        #print(line)


with open('id_number.txt') as f:
    for line in f:
        roll_num.append(line)
        
while( int(time.time() - start_time) < capture_duration):

        _,img=cam.read();
        
        gray=cv2.cvtColor(img,cv2.COLOR_BGR2GRAY)
        faces=faceDetect.detectMultiScale(gray,1.3,5);
        for(x,y,w,h) in faces:
                #cv2.rectangle(img,(x,y),(x+w,y+h),(0,0,255),2)
                cv2.rectangle(img,(x-50,y-50),(x+w+50,y+h+50),(255,0,0),2)
                id,conf=rec.predict(gray[y:y+h,x:x+w])
                if (conf<60):
                        writter(name_list_[id], roll_num[id])
                        if id == 1:
                            pass
                            #ser.write(b'1')
                        if id == 2:
                            pass
                            #ser.write(b'2')
                        if id == 3:
                            pass
                            #ser.write(b'3')
                        if id == 4:
                            pass
                            #ser.write(b'4')
                else:
                        pass
                        #print('unknown')
        cv2.imshow("Face",img);
        if(cv2.waitKey(1)==ord('q')):
                break;


path = r"data.xlsx"
book = load_workbook(path)
len_ = len(book.sheetnames)

sheetname = 'day'+ str(int(len_) + 1)

attendance = sent_mail(name_list_, list_name)


#ser.write(b'9') ## to absent

if attendance[0]=='Absent':
    pass
    #ser.write(b'1')

if attendance[1]=='Absent':
    pass
    #ser.write(b'2')


if attendance[2]=='Absent':
    pass
    #ser.write(b'3')



    
df = pd.DataFrame(list(zip(name_list_, roll_num, attendance)),
               columns =['Name', 'ROll Number', 'Attendance'])



writer = pd.ExcelWriter(path, engine = 'openpyxl')
sheetname = 'day'+ str(int(len_) + 1)
writer.book = book
df.to_excel(writer, sheet_name = sheetname)
writer.save()

if len_ % 1 == 0:
    report_send_mail('Attendance','list', 'bhoomikase15@gmail.com')

cam.release()
cv2.destroyAllWindows()


if __name__ == '__main__':
    detect()
    pass

